"""Quick profile of source xlsx files in D:\\Cursor\\kassa\\ (9 активных; Док 2 не используется).

Prints, for each file:
  - sheet names + dimensions
  - first ~10 non-empty rows (raw, to see structure / headers / context)
  - last 5 non-empty rows
"""
import os
import openpyxl

ROOT = r"D:\Cursor\kassa"
EXCEL_REVIEWED = os.path.join(ROOT, "data", "excel-reviewed")


def xlsx_path(filename: str) -> str:
    p = os.path.join(EXCEL_REVIEWED, filename)
    return p if os.path.isfile(p) else os.path.join(ROOT, filename)


DOCS = [
    ("Док 1", r"client_ведомость с кассами новая база.xlsx"),
    ("Док 3", r"client_ведомость с кассами старая влад и фо расходники и приходники.xlsx"),
    ("Док 4", r"client_ведомость полная по всем движениям с кассами новая база.xlsx"),
    ("Док 5", r"client_ведомость с кассами старая_2.xlsx"),
    ("Док 6", r"client_приходы банк за всё время.xlsx"),
    ("Док 7", r"Еженедельный отчет 01.01.2025 - 31.12.2025_4037299.xlsx"),
    ("Док 8", r"Сверка реализаций по wb 24-25год.xlsx"),
    ("Док 9", r"ИМВРБ главный с аудитом.xlsx"),
    ("Док 10", r"Итог аудит.xlsx"),
]


def fmt_row(vals, max_w=60):
    out = []
    for v in vals:
        if v is None:
            out.append("·")
        else:
            s = str(v)
            if len(s) > max_w:
                s = s[: max_w - 1] + "…"
            out.append(s)
    return " | ".join(out)


for doc_id, fname in DOCS:
    path = xlsx_path(fname)
    if not os.path.exists(path):
        print(f"### {doc_id}  MISSING  {fname}")
        continue
    size = os.path.getsize(path)
    print(f"\n############### {doc_id}  ({size:,} bytes)  {fname} ###############")
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        print(f"  ERROR opening: {e}")
        continue
    for sn in wb.sheetnames:
        ws = wb[sn]
        print(f"\n  ----- SHEET '{sn}'  rows={ws.max_row}  cols={ws.max_column} -----")
        # collect non-empty rows
        rows_data = []
        for r in range(1, ws.max_row + 1):
            vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            if any(v not in (None, "") for v in vals):
                rows_data.append((r, vals))
        print(f"  non-empty rows: {len(rows_data)}")
        if not rows_data:
            continue
        # first 10
        print("  ---- HEAD ----")
        for r, vals in rows_data[:10]:
            print(f"  R{r:>5}: {fmt_row(vals)}")
        # last 5
        if len(rows_data) > 10:
            print("  ---- TAIL ----")
            for r, vals in rows_data[-5:]:
                print(f"  R{r:>5}: {fmt_row(vals)}")
